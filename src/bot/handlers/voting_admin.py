from datetime import datetime, timedelta
from html import escape
from io import BytesIO

from aiogram import F, Router
from aiogram.enums import ChatType
from aiogram.filters import Command
from aiogram.filters.command import CommandObject
from aiogram.fsm.context import FSMContext
from aiogram.types import BufferedInputFile, CallbackQuery, Message
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from config import ADMIN_TELEGRAM_IDS, DATETIME_FORMAT, UFA_TZ
from src.bot.keyboards import voting_keyboards
from src.bot.states import voting_states
from src.database import get_session
from src.database.crud import (
    create_voting,
    get_voting,
    list_vote_history_for_export,
    list_votes_for_export,
    list_votings,
    voting_result_counts,
)


voting_admin_router = Router(name="voting_admin")
voting_admin_router.message.filter(
    lambda message: message.chat.type == ChatType.PRIVATE and message.from_user.id in ADMIN_TELEGRAM_IDS
)
voting_admin_router.callback_query.filter(
    lambda call: bool(
        call.message
        and call.message.chat.type == ChatType.PRIVATE
        and call.from_user.id in ADMIN_TELEGRAM_IDS
    )
)


def _parse_datetime(value: str) -> datetime:
    return datetime.strptime(value.strip(), DATETIME_FORMAT).replace(tzinfo=UFA_TZ)


def _format_datetime(value: datetime) -> str:
    if value.tzinfo is None:
        value = value.replace(tzinfo=UFA_TZ)
    return value.astimezone(UFA_TZ).strftime(DATETIME_FORMAT)


def _parse_voting_id(command: CommandObject | None) -> int | None:
    raw = (command.args or "").strip() if command else ""
    try:
        value = int(raw)
    except (TypeError, ValueError):
        return None
    return value if value > 0 else None


def _status(voting) -> str:
    now = datetime.now(tz=UFA_TZ)
    starts_at = voting.starts_at if voting.starts_at.tzinfo else voting.starts_at.replace(tzinfo=UFA_TZ)
    ends_at = voting.ends_at if voting.ends_at.tzinfo else voting.ends_at.replace(tzinfo=UFA_TZ)
    starts_at = starts_at.astimezone(UFA_TZ)
    ends_at = ends_at.astimezone(UFA_TZ)
    if not voting.active:
        return "закрыто администратором"
    if now < starts_at:
        return "ожидает начала"
    if now >= ends_at:
        return "завершено"
    if voting.allows_changes(now):
        return "идёт, изменения разрешены"
    return "идёт, голоса зафиксированы"


def _voting_text(voting) -> str:
    candidates = "\n".join(
        f"{candidate.position}. {escape(candidate.name)}"
        for candidate in voting.candidates
    )
    description = f"\n{escape(voting.description)}\n" if voting.description else "\n"
    return (
        f"<b>Голосование #{voting.id}: {escape(voting.name)}</b>\n"
        f"{description}\n"
        f"Статус: <b>{_status(voting)}</b>\n"
        f"Начало: <b>{_format_datetime(voting.starts_at)}</b>\n"
        f"Окончание: <b>{_format_datetime(voting.ends_at)}</b>\n"
        f"Изменения до: <b>{_format_datetime(voting.change_deadline)}</b>\n\n"
        f"<b>Кандидаты:</b>\n{candidates}"
    )


def _parse_candidates(value: str) -> list[tuple[str, str | None]]:
    lines = [line.strip() for line in value.splitlines() if line.strip()]
    if len(lines) != 5:
        raise ValueError("Нужно указать ровно 5 кандидатов")
    candidates: list[tuple[str, str | None]] = []
    for line in lines:
        name, separator, description = line.partition("|")
        name = name.strip()
        description = description.strip() if separator else ""
        if not name or len(name) > 255:
            raise ValueError("Имя кандидата должно содержать от 1 до 255 символов")
        candidates.append((name, description or None))
    normalized_names = {name.casefold() for name, _ in candidates}
    if len(normalized_names) != len(candidates):
        raise ValueError("Имена кандидатов не должны повторяться")
    return candidates


async def _send_votings(message: Message, *, edit: bool = False) -> None:
    async with get_session() as session:
        votings = await list_votings(session)
    text = "<b>Голосования</b>\n\nВыберите голосование:" if votings else "Голосований пока нет. Используйте /create_voting."
    markup = voting_keyboards.admin_votings(votings) if votings else None
    if edit:
        await message.edit_text(text, reply_markup=markup)
    else:
        await message.answer(text, reply_markup=markup)


async def _links_text(bot, voting) -> str:
    me = await bot.get_me()
    lines = [f"<b>Ссылки для голосования #{voting.id}</b>"]
    for candidate in voting.candidates:
        if candidate.active:
            url = f"https://t.me/{me.username}?start=vote_{voting.id}_{candidate.id}"
            lines.append(f"\n<b>{candidate.position}. {escape(candidate.name)}</b>\n<code>{url}</code>")
    return "\n".join(lines)


async def _results_text(session, voting) -> str:
    counts = await voting_result_counts(session, voting.id)
    total = sum(count for _, count in counts)
    lines = [f"<b>Результаты: {escape(voting.name)}</b>", ""]
    for position, (candidate, count) in enumerate(counts, start=1):
        lines.append(f"{position}. {escape(candidate.name)} — <b>{count}</b>")
    lines.extend(["", f"Всего проголосовало: <b>{total}</b>"])
    return "\n".join(lines)


def _prepare_sheet(ws, headers: list[str]) -> None:
    ws.append(headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    fill = PatternFill(fill_type="solid", start_color="D9E1F2", end_color="D9E1F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill


def _autosize(ws) -> None:
    for cells in ws.columns:
        width = max(len(str(cell.value or "")) for cell in cells)
        ws.column_dimensions[cells[0].column_letter].width = min(max(width + 2, 10), 70)


async def _build_export(session, voting) -> bytes:
    votes = await list_votes_for_export(session, voting.id)
    history = await list_vote_history_for_export(session, voting.id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Голоса"
    _prepare_sheet(ws, ["Telegram ID", "Имя", "Кандидат", "Причина", "Первый голос (Уфа)", "Обновлено (Уфа)", "Изменений"])
    for vote in votes:
        full_name = " ".join(filter(None, [vote.user.first_name, vote.user.last_name])).strip()
        ws.append([
            vote.user_id,
            full_name,
            vote.candidate.name,
            vote.reason,
            _format_datetime(vote.created_at),
            _format_datetime(vote.updated_at),
            vote.change_count,
        ])
    _autosize(ws)

    history_ws = wb.create_sheet("История изменений")
    _prepare_sheet(history_ws, ["Telegram ID", "Было", "Стало", "Старая причина", "Новая причина", "Изменено (Уфа)"])
    for item in history:
        history_ws.append([
            item.vote.user_id,
            item.previous_candidate.name,
            item.new_candidate.name,
            item.previous_reason,
            item.new_reason,
            _format_datetime(item.created_at),
        ])
    _autosize(history_ws)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@voting_admin_router.message(Command("create_voting"))
async def create_voting_command(message: Message, state: FSMContext):
    await state.clear()
    await state.set_state(voting_states.CreateVoting.name)
    await message.answer("Введите название голосования.")


@voting_admin_router.message(Command("votings"))
async def votings_command(message: Message, state: FSMContext):
    await state.clear()
    await _send_votings(message)


@voting_admin_router.message(Command("voting_results"))
async def voting_results_command(message: Message, command: CommandObject):
    voting_id = _parse_voting_id(command)
    if voting_id is None:
        return await message.answer("Использование: <code>/voting_results ID</code>")
    async with get_session() as session:
        voting = await get_voting(session, voting_id)
        if voting is None:
            return await message.answer("Голосование не найдено.")
        text = await _results_text(session, voting)
    await message.answer(text)


@voting_admin_router.message(Command("voting_links"))
async def voting_links_command(message: Message, command: CommandObject):
    voting_id = _parse_voting_id(command)
    if voting_id is None:
        return await message.answer("Использование: <code>/voting_links ID</code>")
    async with get_session() as session:
        voting = await get_voting(session, voting_id)
    if voting is None:
        return await message.answer("Голосование не найдено.")
    await message.answer(await _links_text(message.bot, voting))


@voting_admin_router.message(Command("cancel"))
async def cancel_voting_flow(message: Message, state: FSMContext):
    await state.clear()
    await message.answer("Действие отменено.")


@voting_admin_router.message(voting_states.CreateVoting.name, F.text)
async def create_voting_name(message: Message, state: FSMContext):
    name = (message.text or "").strip()
    if not 1 <= len(name) <= 255:
        return await message.answer("Название должно содержать от 1 до 255 символов.")
    await state.update_data(name=name)
    await state.set_state(voting_states.CreateVoting.description)
    await message.answer("Введите описание голосования или отправьте <code>-</code>, чтобы пропустить.")


@voting_admin_router.message(voting_states.CreateVoting.description, F.text)
async def create_voting_description(message: Message, state: FSMContext):
    raw = (message.text or "").strip()
    await state.update_data(description=None if raw == "-" else raw)
    await state.set_state(voting_states.CreateVoting.starts_at)
    await message.answer(f"Введите дату и время начала в формате <code>{datetime.now(tz=UFA_TZ).strftime(DATETIME_FORMAT)}</code>.")


@voting_admin_router.message(voting_states.CreateVoting.starts_at, F.text)
async def create_voting_starts_at(message: Message, state: FSMContext):
    try:
        starts_at = _parse_datetime(message.text or "")
    except ValueError:
        return await message.answer("Неверный формат. Пример: <code>15.09.2026 12:00</code>.")
    await state.update_data(starts_at=starts_at)
    await state.set_state(voting_states.CreateVoting.ends_at)
    await message.answer("Введите дату и время окончания. Пример: <code>30.09.2026 20:00</code>.")


@voting_admin_router.message(voting_states.CreateVoting.ends_at, F.text)
async def create_voting_ends_at(message: Message, state: FSMContext):
    try:
        ends_at = _parse_datetime(message.text or "")
    except ValueError:
        return await message.answer("Неверный формат. Пример: <code>30.09.2026 20:00</code>.")
    data = await state.get_data()
    starts_at = data.get("starts_at")
    if not isinstance(starts_at, datetime) or ends_at <= starts_at:
        return await message.answer("Окончание должно быть позже начала голосования.")
    if ends_at - starts_at < timedelta(days=3):
        return await message.answer("Голосование должно длиться не менее трёх суток.")
    await state.update_data(ends_at=ends_at)
    await state.set_state(voting_states.CreateVoting.candidates)
    await message.answer(
        "Отправьте ровно 5 кандидатов — каждый с новой строки.\n"
        "Можно добавить описание после символа <code>|</code>.\n\n"
        "Пример:\n<code>Маруна | Всегда помогает участникам\nИся\nАнна\nОлег\nСаша</code>"
    )


@voting_admin_router.message(voting_states.CreateVoting.candidates, F.text)
async def create_voting_candidates(message: Message, state: FSMContext):
    try:
        candidates = _parse_candidates(message.text or "")
    except ValueError as exc:
        return await message.answer(str(exc))
    data = await state.get_data()
    async with get_session() as session:
        voting = await create_voting(
            session,
            name=data["name"],
            description=data.get("description"),
            starts_at=data["starts_at"],
            ends_at=data["ends_at"],
            candidates=candidates,
        )
    await state.clear()
    await message.answer(
        f"✅ Голосование #{voting.id} создано.\n\n{_voting_text(voting)}",
        reply_markup=voting_keyboards.admin_voting(voting),
    )


@voting_admin_router.callback_query(F.data.startswith("vadmin:"))
async def voting_admin_callback(call: CallbackQuery, state: FSMContext):
    data = (call.data or "").split(":")
    action = data[1] if len(data) > 1 else ""
    if action == "list":
        await state.clear()
        await _send_votings(call.message, edit=True)
        return await call.answer()
    if len(data) != 3 or not data[2].isdigit():
        return await call.answer()
    voting_id = int(data[2])
    async with get_session() as session:
        voting = await get_voting(session, voting_id)
        if voting is None:
            return await call.answer("Голосование не найдено.", show_alert=True)

        if action in {"open", "close"}:
            voting.active = action == "open"
            await session.commit()
            await session.refresh(voting)
            voting = await get_voting(session, voting_id)
            await call.message.edit_text(_voting_text(voting), reply_markup=voting_keyboards.admin_voting(voting))
            return await call.answer("Статус обновлён.")
        if action == "results":
            text = await _results_text(session, voting)
            await call.message.answer(text)
            return await call.answer()
        if action == "export":
            content = await _build_export(session, voting)
            filename = f"voting_{voting.id}_{datetime.now(tz=UFA_TZ):%Y-%m-%d}.xlsx"
            await call.message.answer_document(BufferedInputFile(content, filename=filename))
            return await call.answer("Excel-файл отправлен.")

    if action == "links":
        await call.message.answer(await _links_text(call.bot, voting))
        return await call.answer()
    if action == "view":
        await state.clear()
        await call.message.edit_text(_voting_text(voting), reply_markup=voting_keyboards.admin_voting(voting))
        return await call.answer()
    return await call.answer()
